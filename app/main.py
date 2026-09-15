from __future__ import annotations

import asyncio
import json
import math
import os
import tempfile
import threading
import uuid
from datetime import datetime
from pathlib import Path
from typing import Callable

from docx import Document
from docx.shared import Pt
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import HTMLResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from .diarization import label_segments
from .paths import ensure_app_data_dir, runtime_uploads_dir
from .youtube import download_audio, is_youtube_url

ROOT = Path(__file__).resolve().parent.parent
STATIC = ROOT / "static"
APP_DATA_DIR = ensure_app_data_dir()
RUNTIME_UPLOADS = runtime_uploads_dir()
LIBRARY_PATH = APP_DATA_DIR / "library.json"
MODEL_STATE_PATH = APP_DATA_DIR / "model-state.json"
LIBRARY_LOCK = threading.Lock()

ALLOWED_EXTENSIONS = {
    "aac", "avi", "flac", "m4a", "mkv", "mov", "mp3", "mp4", "mpeg", "mpg",
    "mts", "ogg", "opus", "ts", "wav", "webm", "wmv",
}
MAX_FILE_BYTES = 750 * 1024 * 1024
WHISPER_MODEL = os.getenv("SONOTYPE_WHISPER_MODEL", "small.en")

app = FastAPI(title="Sonotype", version="1.0.0")
app.mount("/static", StaticFiles(directory=STATIC), name="static")


@app.get("/", response_class=HTMLResponse)
def index() -> str:
    return (STATIC / "index.html").read_text(encoding="utf-8")


@app.get("/health")
def health() -> dict:
    return {
        "status": "ok",
        "mode": "local-only",
        "whisper_model": WHISPER_MODEL,
        "named_speaker_recognition": False,
        "diarization_model_configured": bool(os.getenv("SONOTYPE_DIARIZATION_MODEL")),
    }


def _read_library() -> list[dict]:
    try:
        entries = json.loads(LIBRARY_PATH.read_text(encoding="utf-8"))
    except (FileNotFoundError, OSError, json.JSONDecodeError):
        return []
    return entries if isinstance(entries, list) else []


def _normalise_library(entries: list[dict]) -> list[dict]:
    cleaned: list[dict] = []
    for raw in entries:
        if not isinstance(raw, dict):
            continue
        text = str(raw.get("text", "")).strip()
        if not text:
            continue
        cleaned.append({
            "id": str(raw.get("id") or uuid.uuid4()),
            "name": str(raw.get("name") or "Sonotype transcript")[:240],
            "format": str(raw.get("format") or "TXT")[:16].upper(),
            "text": text,
            "confidence": str(raw.get("confidence") or "")[:80],
            "savedAt": str(raw.get("savedAt") or datetime.now().isoformat()),
        })
        if len(cleaned) == 40:
            break
    return cleaned


def _write_library(entries: list[dict]) -> None:
    payload = json.dumps(_normalise_library(entries), ensure_ascii=False, indent=2)
    temporary = LIBRARY_PATH.with_suffix(".tmp")
    temporary.write_text(payload, encoding="utf-8")
    temporary.replace(LIBRARY_PATH)


def _model_is_ready() -> bool:
    model_path = Path(WHISPER_MODEL).expanduser()
    if model_path.is_dir():
        return True

    try:
        state = json.loads(MODEL_STATE_PATH.read_text(encoding="utf-8"))
    except (FileNotFoundError, OSError, json.JSONDecodeError):
        state = {}
    if state.get("model") == WHISPER_MODEL and state.get("ready") is True:
        return True

    # Recognise a model downloaded by an earlier Sonotype release so an
    # upgrade does not show the first-run prompt or download it again.
    cache_root = os.getenv("HF_HUB_CACHE") or os.getenv("HUGGINGFACE_HUB_CACHE")
    if not cache_root:
        hf_home = os.getenv("HF_HOME")
        cache_root = str(Path(hf_home) / "hub") if hf_home else str(Path.home() / ".cache" / "huggingface" / "hub")
    model_cache = Path(cache_root) / f"models--Systran--faster-whisper-{WHISPER_MODEL.replace('/', '--')}"
    return any(model_cache.glob("snapshots/*/model.bin"))


def _mark_model_ready() -> None:
    temporary = MODEL_STATE_PATH.with_suffix(".tmp")
    temporary.write_text(
        json.dumps({"model": WHISPER_MODEL, "ready": True}, indent=2),
        encoding="utf-8",
    )
    temporary.replace(MODEL_STATE_PATH)


@app.get("/runtime")
def runtime_status() -> dict:
    """Expose only local readiness state needed by the desktop UI."""
    with LIBRARY_LOCK:
        library_count = len(_read_library())
    return {"model_ready": _model_is_ready(), "library_count": library_count}


@app.get("/library")
def get_library() -> list[dict]:
    with LIBRARY_LOCK:
        return _read_library()


@app.put("/library")
async def put_library(entries: list[dict]) -> list[dict]:
    cleaned = _normalise_library(entries)
    with LIBRARY_LOCK:
        _write_library(cleaned)
    return cleaned


def _format_time(seconds: float) -> str:
    seconds = max(0, int(seconds))
    return f"{seconds // 3600:02}:{(seconds % 3600) // 60:02}:{seconds % 60:02}"


def _render_transcript(segments: list[dict]) -> str:
    lines = []
    for segment in segments:
        prefix = f"[{_format_time(segment['start'])}] "
        if segment.get("speaker"):
            prefix += f"{segment['speaker']}: "
        lines.append(prefix + segment["text"].strip())
    return "\n".join(line for line in lines if line.strip())


def _transcribe(audio_path: Path, diarize: bool, send: Callable[[dict], None]) -> dict:
    try:
        from faster_whisper import WhisperModel
        from faster_whisper.audio import decode_audio
    except ImportError as exc:
        raise RuntimeError("faster-whisper is not installed. Install requirements.txt first.") from exc

    send({"type": "phase", "phase": "Loading the local transcription model…"})
    audio = decode_audio(str(audio_path), sampling_rate=16000)
    duration = len(audio) / 16000 if len(audio) else 0
    model = WhisperModel(WHISPER_MODEL, device="cpu", compute_type="int8")
    send({"type": "phase", "phase": "Transcribing locally on CPU…"})
    segment_stream, info = model.transcribe(
        audio, beam_size=5, vad_filter=True, word_timestamps=False
    )

    segments: list[dict] = []
    scores: list[float] = []
    last_progress = -1
    for item in segment_stream:
        segments.append({"start": round(item.start, 2), "end": round(item.end, 2), "text": item.text.strip()})
        if item.avg_logprob is not None:
            scores.append(math.exp(min(0.0, item.avg_logprob)))
        progress = min(99, int((item.end / duration) * 100)) if duration else 0
        if progress > last_progress:
            last_progress = progress
            send({"type": "progress", "pct": progress})

    with LIBRARY_LOCK:
        _mark_model_ready()

    if diarize:
        try:
            send({"type": "phase", "phase": "Assigning anonymous speaker labels…"})
            segments = label_segments(segments, audio_path)
        except RuntimeError:
            # Diarization model not configured — skip speaker labels, keep transcript
            pass

    confidence = round((sum(scores) / len(scores) * 100) if scores else 0, 1)
    return {
        "transcript": _render_transcript(segments), "segments": segments,
        "confidence": confidence, "duration_seconds": round(duration, 1),
        "language": getattr(info, "language", "unknown"),
    }


def _stream_job(job: Callable[[Callable[[dict], None]], dict]) -> StreamingResponse:
    queue: asyncio.Queue[dict] = asyncio.Queue()
    loop = asyncio.get_running_loop()

    def send(event: dict) -> None:
        loop.call_soon_threadsafe(queue.put_nowait, event)

    def run() -> None:
        try:
            result = job(send)
            send({"type": "progress", "pct": 100})
            send({"type": "done", **result})
        except Exception as exc:
            send({"type": "error", "message": str(exc)})

    async def event_stream():
        loop.run_in_executor(None, run)
        while True:
            event = await queue.get()
            yield f"data: {json.dumps(event)}\n\n"
            if event["type"] in {"done", "error"}:
                break

    return StreamingResponse(event_stream(), media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.post("/transcribe")
async def transcribe(media: UploadFile = File(...), diarize: bool = Form(default=True)) -> StreamingResponse:
    extension = (media.filename or "").rsplit(".", 1)[-1].lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail="Unsupported audio or video file type.")

    with tempfile.NamedTemporaryFile(dir=RUNTIME_UPLOADS, suffix=f".{extension}", delete=False) as destination:
        temp_path = Path(destination.name)
        total = 0
        while chunk := await media.read(1024 * 1024):
            total += len(chunk)
            if total > MAX_FILE_BYTES:
                temp_path.unlink(missing_ok=True)
                raise HTTPException(status_code=413, detail="File exceeds the 750 MB local limit.")
            destination.write(chunk)

    def job(send: Callable[[dict], None]) -> dict:
        try:
            return _transcribe(temp_path, diarize, send)
        finally:
            temp_path.unlink(missing_ok=True)

    return _stream_job(job)


@app.post("/transcribe/youtube")
async def transcribe_youtube(url: str = Form(...), diarize: bool = Form(default=True)) -> StreamingResponse:
    if not is_youtube_url(url):
        raise HTTPException(status_code=400, detail="Enter a valid YouTube URL.")

    def job(send: Callable[[dict], None]) -> dict:
        temp_path: Path | None = None
        try:
            send({"type": "phase", "phase": "Downloading audio locally…"})

            def download_progress(status: dict) -> None:
                if status.get("status") == "downloading":
                    total = status.get("total_bytes") or status.get("total_bytes_estimate")
                    downloaded = status.get("downloaded_bytes", 0)
                    if total:
                        send({"type": "progress", "pct": min(25, int(downloaded / total * 25))})

            temp_path = download_audio(url, RUNTIME_UPLOADS, download_progress)
            send({"type": "phase", "phase": "Download complete. Preparing local transcription…"})
            return _transcribe(temp_path, diarize, send)
        finally:
            if temp_path:
                temp_path.unlink(missing_ok=True)

    return _stream_job(job)


def _safe_filename(kind: str) -> str:
    return f"sonotype_transcript_{datetime.now():%Y%m%d_%H%M%S}.{kind}"


def _validate_transcript(transcript: str) -> str:
    cleaned = transcript.strip()
    if not cleaned:
        raise HTTPException(status_code=400, detail="A transcript is required.")
    return cleaned


@app.post("/export/docx")
async def export_docx(transcript: str = Form(...)) -> Response:
    text = _validate_transcript(transcript)
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    document.add_heading("Sonotype Transcript", level=0)
    for line in text.splitlines():
        document.add_paragraph(line)
    buffer = tempfile.SpooledTemporaryFile()
    document.save(buffer)
    buffer.seek(0)
    return Response(content=buffer.read(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f'attachment; filename="{_safe_filename("docx")}"'})


@app.post("/export/pdf")
async def export_pdf(transcript: str = Form(...)) -> Response:
    from io import BytesIO
    from xml.sax.saxutils import escape

    text = _validate_transcript(transcript)
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=0.75 * inch, rightMargin=0.75 * inch)
    styles = getSampleStyleSheet()
    story = [Paragraph("Sonotype Transcript", styles["Title"]), Spacer(1, 0.2 * inch)]
    for line in text.splitlines():
        story.append(Paragraph(escape(line) or " ", styles["BodyText"]))
        story.append(Spacer(1, 0.06 * inch))
    document.build(story)
    return Response(content=buffer.getvalue(), media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{_safe_filename("pdf")}"'})
