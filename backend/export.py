"""Generic document exports for locally produced transcripts."""

import io
from datetime import datetime
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _metadata(document: Document, source_name: str, confidence: str, duration: str) -> None:
    metadata = document.add_paragraph()
    metadata.add_run("Source: ").bold = True
    metadata.add_run(source_name or "Local recording")
    metadata.add_run("\nGenerated: ").bold = True
    metadata.add_run(datetime.now().astimezone().strftime("%Y-%m-%d %H:%M %Z"))
    if confidence:
        metadata.add_run("\nEstimated confidence: ").bold = True
        metadata.add_run(f"{confidence}%")
    if duration:
        metadata.add_run("\nRecording duration: ").bold = True
        metadata.add_run(duration)


def generate_docx(transcript: str, source_name: str = "", confidence: str = "", duration: str = "") -> bytes:
    document = Document()
    document.add_heading("Transcript", level=1)
    _metadata(document, source_name, confidence, duration)
    document.add_paragraph()
    for paragraph in transcript.split("\n\n"):
        document.add_paragraph(paragraph.strip())
    footer = document.sections[0].footer.paragraphs[0]
    footer.text = "Created locally"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def generate_pdf(transcript: str, source_name: str = "", confidence: str = "", duration: str = "") -> bytes:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    content = [Paragraph("Transcript", styles["Title"])]
    detail = f"Source: {escape(source_name or 'Local recording')}<br/>Generated: {datetime.now().astimezone().strftime('%Y-%m-%d %H:%M %Z')}"
    if confidence:
        detail += f"<br/>Estimated confidence: {confidence}%"
    if duration:
        detail += f"<br/>Recording duration: {duration}"
    content.extend([Paragraph(detail, styles["Normal"]), Spacer(1, 16)])
    for paragraph in transcript.split("\n\n"):
        content.extend([Paragraph(escape(paragraph).replace("\n", "<br/>"), styles["BodyText"]), Spacer(1, 8)])
    SimpleDocTemplate(buffer, pagesize=LETTER).build(content)
    return buffer.getvalue()
